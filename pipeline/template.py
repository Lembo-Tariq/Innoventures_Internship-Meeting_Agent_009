# pipeline/template.py
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.table import WD_TABLE_ALIGNMENT
from models.meeting import MeetingData

TEMPLATE_PATH = os.path.join(os.path.dirname(os.path.dirname(__file__)), "templates", "mom_template_v3.docx")
OUTPUT_DIR    = os.path.join(os.path.dirname(os.path.dirname(__file__)), "output")


def populate_mom_template(data: MeetingData) -> str:
    doc = Document(TEMPLATE_PATH)

    _replace_all(doc, {
        "{{TITLE}}":     data.title,
        "{{DATE}}":      str(data.date),
        "{{ATTENDEES}}": "\n".join(f"• {a}" for a in data.attendees),
    })

    for placeholder, items in [
        ("{{AGENDA}}",            data.agenda),
        ("{{DISCUSSION_POINTS}}", data.discussion_points),
        ("{{DECISIONS}}",         data.decisions),
        ("{{FOLLOW_UPS}}",        data.follow_ups),
    ]:
        _replace_with_bullets(doc, placeholder, items)

    _rebuild_action_table(doc, data)
    _fix_pagination(doc)

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    safe = data.title.replace("/", "-").replace("\\", "-").replace(":", "").replace('"', "").strip()[:80]
    path = os.path.join(OUTPUT_DIR, f"MoM - {safe} - {data.date}.docx")
    doc.save(path)
    return path


def _replace_all(doc, replacements):
    for para in doc.paragraphs:
        _replace_para(para, replacements)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    _replace_para(para, replacements)


def _replace_para(para, replacements):
    for key, val in replacements.items():
        if key in para.text:
            for run in para.runs:
                if key in run.text:
                    run.text = run.text.replace(key, val)


def _replace_with_bullets(doc, placeholder, items):
    for para in doc.paragraphs:
        if placeholder not in para.text:
            continue
        para.text = ""
        if not items:
            para.text = "None recorded."
            return
        para.text = f"• {items[0]}"
        para.style = doc.styles["Normal"]
        para.paragraph_format.keep_together = True
        for item in items[1:]:
            new_para = doc.add_paragraph(f"• {item}")
            new_para.style = doc.styles["Normal"]
            new_para.paragraph_format.keep_together = True
            para._element.addnext(new_para._element)
            para = new_para
        return


def _rebuild_action_table(doc, data: MeetingData):
    """
    The template has a placeholder table with a dummy {{ACTION_ITEMS}} row.
    We delete that table and replace it with a real one containing actual data.
    Then we re-insert it after the Action Items heading paragraph.
    """
    for table in doc.tables:
        if not table.rows:
            continue
        header_cells = [c.text.strip() for c in table.rows[0].cells]
        if "Task" not in header_cells:
            continue

        # Delete the placeholder table
        table._element.getparent().remove(table._element)

        # Build the real table
        new_table = doc.add_table(rows=1, cols=2)
        new_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _apply_table_borders(new_table)

        # Header row
        for i, label in enumerate(["Task", "Assignee"]):
            cell = new_table.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(label)
            run.bold = True
            run.font.size = Pt(11)
            run.font.name = "Calibri"
            run.font.color.rgb = RGBColor(255, 255, 255)
            _shade_cell(cell, "0A1628")
            cell.paragraphs[0].paragraph_format.keep_with_next = True

        new_table.rows[0].cells[0].width = Inches(4.5)
        new_table.rows[0].cells[1].width = Inches(2.0)

        # Data rows
        for item in data.action_items:
            row = new_table.add_row()
            row.cells[0].text = item.task
            row.cells[1].text = item.assignee
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.keep_together = True
                    for run in para.runs:
                        run.font.size = Pt(10)
                        run.font.name = "Calibri"

        # Insert table after "Action Items" heading
        for para in doc.paragraphs:
            if "Action Items" in para.text:
                para._element.addnext(new_table._element)
                break
        return


def _apply_table_borders(table):
    tbl = table._element
    tblPr = tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    tblBorders = OxmlElement("w:tblBorders")
    for name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = OxmlElement(f"w:{name}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "D0D9E8")
        tblBorders.append(el)
    tblPr.append(tblBorders)


def _shade_cell(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    tcPr.append(tcPr.makeelement(qn("w:shd"), {
        qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): hex_color
    }))


def _fix_pagination(doc):
    for para in doc.paragraphs:
        pf = para.paragraph_format
        pf.widow_control = True
        if para.runs and para.runs[0].bold:
            pf.keep_with_next = True
            pf.space_before = Pt(10)
            pf.space_after = Pt(4)
        elif para.text.startswith("•"):
            pf.keep_together = True
            pf.keep_with_next = False
            pf.space_before = Pt(2)
            pf.space_after = Pt(2)
        elif not para.text.strip():
            pf.space_before = Pt(0)
            pf.space_after = Pt(2)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.keep_together = True

    prev_empty = False
    to_remove = []
    for para in doc.paragraphs:
        empty = not para.text.strip()
        if empty and prev_empty:
            to_remove.append(para)
        prev_empty = empty
    for para in to_remove:
        para._element.getparent().remove(para._element)
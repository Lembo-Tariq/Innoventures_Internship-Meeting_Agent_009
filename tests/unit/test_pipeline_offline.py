# tests/unit/test_pipeline_offline.py
"""Offline tests: the LLM is stubbed, so these run without an API key.

Run with:  pytest
"""
import asyncio
import json
import os
import sys
import types

import pytest
from docx import Document

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))

import pipeline.extract as extract
import pipeline.template as template
from models.meeting import MeetingData

VALID = {
    "title": "Q3 Planning Sync",
    "date": "2026-05-18",
    "attendees": ["Sarah (Designer) - Marketing", "Omar (IT) - InnoVentures"],
    "agenda": ["Onboarding flow", "Meeting notes automation"],
    "discussion_points": ["Wireframes need one more revision for mobile."],
    "decisions": ["Prioritise the meeting agent PoC"],
    "action_items": [
        {"task": "Revise wireframes and hand off to dev", "assignee": "Sarah (Designer) - Marketing",
         "deadline": "2026-05-21", "priority": "high"},
        {"task": "Build the PoC", "assignee": "Omar (IT) - InnoVentures"},
    ],
    "follow_ups": ["Check PoC status after two weeks"],
}


class FakeClient:
    """Returns the queued responses in order and records every request."""

    def __init__(self, responses):
        self.responses = list(responses)
        self.calls = []
        self.chat = types.SimpleNamespace(completions=self)

    def create(self, **kwargs):
        self.calls.append([dict(m) for m in kwargs["messages"]])
        content = self.responses.pop(0)
        msg = types.SimpleNamespace(content=content)
        return types.SimpleNamespace(choices=[types.SimpleNamespace(message=msg)])


@pytest.fixture
def fake(monkeypatch):
    def install(*responses):
        client = FakeClient(responses)
        monkeypatch.setattr(extract, "get_client", lambda: client)
        return client
    return install


def run(transcript="Sarah: hi"):
    return asyncio.run(extract.extract_meeting_data(transcript))


def test_valid_json_parses_first_try(fake):
    client = fake(json.dumps(VALID))
    data = run()
    assert data.title == "Q3 Planning Sync"
    assert data.action_items[0].deadline == "2026-05-21"
    assert data.action_items[1].priority == "medium"  # schema default
    assert len(client.calls) == 1


def test_markdown_fenced_json_is_accepted(fake):
    fake("```json\n" + json.dumps(VALID) + "\n```")
    assert run().title == "Q3 Planning Sync"


def test_invalid_output_is_retried_with_error_feedback(fake):
    bad = json.dumps({**VALID, "date": "next Tuesday"})
    client = fake(bad, json.dumps(VALID))
    data = run()
    assert str(data.date) == "2026-05-18"
    assert len(client.calls) == 2
    # The second request must show the model its bad answer and the validation error.
    retry_messages = client.calls[1]
    assert retry_messages[-2] == {"role": "assistant", "content": bad}
    assert "date" in retry_messages[-1]["content"]


def test_non_json_is_retried(fake):
    client = fake("Sure! Here are the minutes...", json.dumps(VALID))
    assert run().title == "Q3 Planning Sync"
    assert len(client.calls) == 2


def test_gives_up_after_max_attempts(fake):
    missing_assignee = json.dumps({**VALID, "action_items": [{"task": "x"}]})
    client = fake(*[missing_assignee] * extract.MAX_ATTEMPTS)
    with pytest.raises(RuntimeError, match="after 3 attempts"):
        run()
    assert len(client.calls) == extract.MAX_ATTEMPTS


def test_prompt_uses_todays_date_at_call_time(fake, monkeypatch):
    class FixedDate(extract.date):
        @classmethod
        def today(cls):
            return cls(2030, 1, 2)
    monkeypatch.setattr(extract, "date", FixedDate)
    client = fake(json.dumps(VALID))
    run()
    assert "2030-01-02" in client.calls[0][0]["content"]


def test_template_renders_all_sections(tmp_path, monkeypatch):
    monkeypatch.setattr(template, "OUTPUT_DIR", str(tmp_path))
    path = template.populate_mom_template(MeetingData(**VALID))

    assert os.path.dirname(path) == str(tmp_path)
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "{{" not in text  # every placeholder was filled
    assert "Onboarding flow" in text and "Check PoC status" in text

    action_table = next(t for t in doc.tables if t.rows[0].cells[0].text == "Task")
    rows = [[c.text for c in r.cells] for r in action_table.rows]
    assert rows[0] == ["Task", "Assignee", "Due"]
    assert rows[1][2] == "2026-05-21"
    assert rows[2][2] == "TBD"


def test_template_handles_empty_sections(tmp_path, monkeypatch):
    monkeypatch.setattr(template, "OUTPUT_DIR", str(tmp_path))
    data = MeetingData(**{**VALID, "decisions": [], "follow_ups": []})
    text = "\n".join(p.text for p in Document(template.populate_mom_template(data)).paragraphs)
    assert "None recorded." in text
